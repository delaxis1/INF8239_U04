"""Genera el informe de síntesis de U04.E07 (DOCX + PDF) desde los artefactos del repositorio.

Lee reports/lab12/*/global_metrics.json, fairness_by_group.csv, fairness_intersectional.csv,
comparison.csv, reports/lab11/risk_register_scored.csv y student_inputs/lab12_decision.yml,
así que el informe siempre coincide con la última ejecución de los scripts. No hay ninguna
cifra escrita a mano en este archivo.

Uso:
    uv run python herramientas/generar_informe.py --repo-url https://github.com/USUARIO/INF8239_U04
    uv run python herramientas/generar_informe.py --sin-pdf      # solo el DOCX

Requiere python-docx. El PDF se exporta con Microsoft Word (Windows).
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from datetime import date
from pathlib import Path

import pandas as pd
import yaml
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BASE = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BASE / "src"))
sys.path.insert(0, str(Path(__file__).resolve().parent))
REPORTS = BASE / "reports"
LAB11 = REPORTS / "lab11"
LAB12 = REPORTS / "lab12"
NOMBRE = "100805370_DeLaCruz_U04E07_Gobernanza"

NAVY = RGBColor(0x12, 0x35, 0x5B)
INK = RGBColor(0x1F, 0x1F, 0x1F)
MUTED = RGBColor(0x5F, 0x5E, 0x5A)
ALERT = RGBColor(0x9B, 0x2C, 0x2C)
HEADER_FILL = "12355B"
ZEBRA_FILL = "F2F4F7"
WARN_FILL = "FBEAEA"
BORDER = "C9CED6"
MESES = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto",
         "septiembre", "octubre", "noviembre", "diciembre"]
SOPORTE_MINIMO = 30


def set_cell_fill(cell, hex_fill: str) -> None:
    props = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), hex_fill)
    props.append(shading)


def set_table_borders(table, color: str = BORDER, size: int = 4, inner: bool = True) -> None:
    props = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    edges = ["top", "left", "bottom", "right"] + (["insideH", "insideV"] if inner else [])
    for edge in edges:
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), str(size))
        element.set(qn("w:color"), color)
        borders.append(element)
    props.append(borders)


def remove_table_borders(table) -> None:
    props = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "nil")
        borders.append(element)
    props.append(borders)


def set_col_widths(table, widths_cm) -> None:
    table.autofit = False
    for row in table.rows:
        for cell, width in zip(row.cells, widths_cm):
            cell.width = Cm(width)


def unbreakable(text: str) -> str:
    text = re.sub(r"(\d) %", "\\1\u00a0%", text)
    return re.sub(r"(?<=\d) (?=\d{3}\b)", "\u00a0", text)


def add_inline_markdown(p, text: str, *, size=10.5, bold=False, color=INK) -> None:
    text = unbreakable(text)
    for token in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if not token:
            continue
        run = p.add_run()
        if token.startswith("**"):
            run.text = token[2:-2]
            run.bold = True
        elif token.startswith("`"):
            run.text = token[1:-1]
            run.font.name = "Consolas"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
        else:
            run.text = token
            run.bold = bold
        run.font.size = Pt(size if not token.startswith("`") else size - 1)
        run.font.color.rgb = color


def cell_text(cell, text: str, *, bold=False, color=INK, size=8.5,
              align=WD_ALIGN_PARAGRAPH.LEFT) -> None:
    cell.text = ""
    paragraph_ = cell.paragraphs[0]
    paragraph_.alignment = align
    paragraph_.paragraph_format.space_after = Pt(0)
    paragraph_.paragraph_format.space_before = Pt(0)
    add_inline_markdown(paragraph_, text, size=size, bold=bold, color=color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_hyperlink(paragraph_, url: str, text: str | None = None) -> None:
    rel_id = paragraph_.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1F5FAF")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.append(color)
    props.append(underline)
    run.append(props)
    text_element = OxmlElement("w:t")
    text_element.text = text or url
    text_element.set(qn("xml:space"), "preserve")
    run.append(text_element)
    link.append(run)
    paragraph_._p.append(link)


def add_field(paragraph_, instruction: str) -> None:
    run = paragraph_.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)
    run.font.size = Pt(8)
    run.font.color.rgb = MUTED


def paragraph(doc, text: str = "", *, size=10.5, bold=False, color=INK, after=6,
              align=None, style=None):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    add_inline_markdown(p, text, size=size, bold=bold, color=color)
    return p


def heading(doc, text: str) -> None:
    p = doc.add_paragraph(style="Heading 1")
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.size = Pt(13)
    run.font.bold = True
    run.font.color.rgb = NAVY


def subheading(doc, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.size = Pt(10.5)
    run.font.bold = True
    run.font.color.rgb = NAVY


def bullet(doc, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    add_inline_markdown(p, text, size=10)


def numbered(doc, text: str, indice: int) -> None:
    """Numeración explícita: la lista reinicia en cada sección en lugar de continuar."""
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.left_indent = Cm(0.62)
    p.paragraph_format.first_line_indent = Cm(-0.62)
    add_inline_markdown(p, f"{indice}.  {text}", size=10)


def caption(doc, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    run = p.add_run(unbreakable(text))
    run.italic = True
    run.font.size = Pt(8.5)
    run.font.color.rgb = MUTED


def keep_together(table) -> None:
    for i, fila in enumerate(table.rows):
        tr_pr = fila._tr.get_or_add_trPr()
        tr_pr.append(OxmlElement("w:cantSplit"))
        if i == 0:
            tr_pr.append(OxmlElement("w:tblHeader"))
        for celda in fila.cells:
            for parrafo in celda.paragraphs:
                parrafo.paragraph_format.keep_with_next = True


def data_table(doc, header, rows, widths, align_right_from=1, bold_rows=(),
               warn_rows=(), size=8.5):
    table = doc.add_table(rows=1 + len(rows), cols=len(header))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)
    for j, title in enumerate(header):
        cell = table.rows[0].cells[j]
        set_cell_fill(cell, HEADER_FILL)
        align = WD_ALIGN_PARAGRAPH.RIGHT if j >= align_right_from else WD_ALIGN_PARAGRAPH.LEFT
        cell_text(cell, title, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF), size=size,
                  align=align)
    for i, values in enumerate(rows, start=1):
        warn = (i - 1) in warn_rows
        for j, value in enumerate(values):
            cell = table.rows[i].cells[j]
            if warn:
                set_cell_fill(cell, WARN_FILL)
            elif i % 2 == 0:
                set_cell_fill(cell, ZEBRA_FILL)
            align = WD_ALIGN_PARAGRAPH.RIGHT if j >= align_right_from else WD_ALIGN_PARAGRAPH.LEFT
            cell_text(cell, value, bold=((i - 1) in bold_rows),
                      color=ALERT if warn else INK, size=size, align=align)
    set_col_widths(table, widths)
    keep_together(table)
    return table


def key_value_table(doc, pairs, widths=(4.2, 12.6)):
    table = doc.add_table(rows=len(pairs), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)
    for i, (key, value) in enumerate(pairs):
        set_cell_fill(table.rows[i].cells[0], ZEBRA_FILL)
        cell_text(table.rows[i].cells[0], key, bold=True, size=9)
        cell_text(table.rows[i].cells[1], value, size=9)
    set_col_widths(table, list(widths))
    keep_together(table)
    return table


def fecha_larga(d: date) -> str:
    return f"{d.day} de {MESES[d.month - 1]} de {d.year}"


def leer_evidencia() -> dict:
    from inf8239_u04.data import load_data, temporal_split
    from inf8239_u04.fairness import intersectional
    from inf8239_u04.governance import load_config

    datos = {}
    for modelo in ("logistic", "tree"):
        carpeta = LAB12 / modelo
        datos[modelo] = {
            "global": json.loads((carpeta / "global_metrics.json").read_text(encoding="utf-8")),
            "grupos": pd.read_csv(carpeta / "fairness_by_group.csv", index_col=0),
            "inter": pd.read_csv(carpeta / "fairness_intersectional.csv", index_col=0),
        }
    datos["comparacion"] = pd.read_csv(LAB12 / "comparison.csv")
    datos["umbral"] = json.loads((LAB12 / "umbral_regla_manual.json").read_text(encoding="utf-8"))
    datos["riesgos"] = pd.read_csv(LAB11 / "risk_register_scored.csv")
    datos["decision"] = yaml.safe_load(
        (BASE / "student_inputs" / "lab12_decision.yml").read_text(encoding="utf-8"))
    datos["gates"] = load_config()["gates"]

    df = load_data()
    train, test = temporal_split(df)
    test = test.copy()
    test["inter"] = intersectional(test)
    soporte = test.groupby("inter").target.agg(n="size", positivos="sum")
    soporte["negativos"] = soporte.n - soporte.positivos
    datos["soporte"] = soporte
    datos["particion"] = {
        "train_n": len(train), "test_n": len(test),
        "train_ini": train.received_date.min().date(),
        "train_fin": train.received_date.max().date(),
        "test_ini": test.received_date.min().date(),
        "test_fin": test.received_date.max().date(),
        "total": len(df), "columnas": df.shape[1],
        "grupos": df.protected_group.value_counts().to_dict(),
        "positivos": float(df.target.mean()),
    }
    return datos


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__,
                                     formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--repo-url", help="enlace del repositorio en GitHub")
    parser.add_argument("--sin-pdf", action="store_true", help="no exportar a PDF")
    args = parser.parse_args()
    from informe_contenido import build, export_pdf
    docx_path = build(leer_evidencia(), args.repo_url)
    print("DOCX:", docx_path)
    if not args.sin_pdf:
        print("PDF: ", export_pdf(docx_path))


if __name__ == "__main__":
    main()
